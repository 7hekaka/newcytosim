#!/usr/bin/env python3
from __future__ import annotations

import csv
import math
import shutil
import subprocess
from dataclasses import dataclass
from pathlib import Path

import matplotlib

matplotlib.use("Agg")
import matplotlib.image as mpimg
import matplotlib.patches as patches
import matplotlib.pyplot as plt
import numpy as np
from PIL import Image, ImageDraw, ImageFont, ImageStat
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt
from scipy.ndimage import gaussian_filter


ROOT = Path(__file__).resolve().parents[1]
JOB_ROOT = ROOT / "turnover_continuous_supply_scaled" / "job12" / "save"
ANALYSIS = ROOT / "analysis" / "results" / "turnover_continuous_supply_scaled_job12_chewer_count2x_rate2x"
OUT = ROOT / "analysis" / "results" / "turnover_aggressive_large_visual_deck"
RAW = OUT / "raw_play"
RAW_ORIENTED = OUT / "raw_oriented"
PROCESSED = OUT / "processed"
FRAMES = OUT / "frames"
PANELS = OUT / "panels"
PLOTS = OUT / "plots"
ASSETS = OUT / "assets"
REFERENCE_BACKGROUND_SCREENSHOT = Path("/mnt/c/Users/User/AppData/Local/Temp/codex-clipboard-ce660a4c-128f-48d7-abb9-98d64ad4edc5.png")

PLAY = JOB_ROOT / "play"
FRAME_INTERVAL_MIN = 0.05
ONSET_MIN = 1.0
FINAL_MIN = 8.0
FINAL_POST_MIN = FINAL_MIN - ONSET_MIN
RENDER_SIZE = 1800
SAMPLES = 8
SMOOTH_FRAME_COUNT = 61
MOVIE_PANEL_W = 455
MOVIE_PANEL_H = 900

RUNS = {
    "nomotor_xlink": {
        "label": "No motors",
        "short": "No motors",
        "run": JOB_ROOT / "r0003",
        "color": "#4d4d4d",
    },
    "rotatable_xlink": {
        "label": "Motors",
        "short": "Motors",
        "run": JOB_ROOT / "r0007",
        "color": "#0a9f7a",
    },
}

PANEL_TIMEPOINTS = [
    ("t00", 0.0, "0 min"),
    ("t02", 2.0, "2 min"),
    ("t04", 4.0, "4 min"),
    ("t06", 6.0, "6 min"),
    ("t07", 7.0, "7 min"),
]

PHASES = [
    ("early", "0-2 min", 0.0, 2.0),
    ("middle", "2-5 min", 2.0, 5.0),
    ("late", "5-7 min", 5.0, FINAL_POST_MIN + 1e-9),
]


@dataclass(frozen=True)
class Camera:
    key: str
    rotation: str
    zoom: float
    view_scale: float
    postprocess_rot90: int = 0


def quat(axis: str, degrees: float) -> tuple[float, float, float, float]:
    angle = math.radians(degrees) / 2.0
    c = math.cos(angle)
    s = math.sin(angle)
    if axis == "x":
        return c, s, 0.0, 0.0
    if axis == "y":
        return c, 0.0, s, 0.0
    if axis == "z":
        return c, 0.0, 0.0, s
    raise KeyError(axis)


def qstr(q: tuple[float, float, float, float]) -> str:
    norm = math.sqrt(sum(value * value for value in q))
    return " ".join(f"{value / norm:.8f}" for value in q)


CAMERA = Camera(
    key="upright_side",
    rotation=qstr(quat("y", 90)),
    zoom=0.92,
    view_scale=38.0,
    postprocess_rot90=1,
)


def post_min_to_frame(post_min: float) -> int:
    global_min = ONSET_MIN + post_min
    return int(round(global_min / FRAME_INTERVAL_MIN))


def f(value: object, default: float = float("nan")) -> float:
    try:
        if value in ("", None):
            return default
        return float(value)
    except (TypeError, ValueError):
        return default


def sem(values: list[float]) -> float:
    arr = np.asarray([value for value in values if np.isfinite(value)], dtype=float)
    if arr.size == 0:
        return float("nan")
    if arr.size == 1:
        return 0.0
    return float(np.nanstd(arr, ddof=1) / math.sqrt(arr.size))


def read_csv(path: Path) -> list[dict[str, str]]:
    with path.open(newline="", encoding="utf-8") as handle:
        return list(csv.DictReader(handle))


def write_csv(path: Path, rows: list[dict[str, object]]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    if not rows:
        path.write_text("", encoding="utf-8")
        return
    fieldnames: list[str] = []
    for row in rows:
        for key in row:
            if key not in fieldnames:
                fieldnames.append(key)
    with path.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.DictWriter(handle, fieldnames=fieldnames)
        writer.writeheader()
        writer.writerows(rows)


def reset_generated_outputs() -> None:
    for directory in (RAW_ORIENTED, PROCESSED, FRAMES, PANELS, PLOTS, ASSETS):
        if directory.exists():
            shutil.rmtree(directory)


def raw_path(condition: str, frame: int, slug: str) -> Path:
    return RAW / f"{condition}_{CAMERA.key}_{slug}_frame{frame:04d}_raw.png"


def raw_oriented_path(condition: str, frame: int, slug: str) -> Path:
    return RAW_ORIENTED / f"{condition}_{CAMERA.key}_{slug}_frame{frame:04d}_raw_oriented.png"


def processed_path(condition: str, frame: int, slug: str) -> Path:
    return PROCESSED / f"{condition}_{CAMERA.key}_{slug}_frame{frame:04d}_cyan.png"


def render_raw(condition: str, frame: int, slug: str) -> tuple[Path, str, int]:
    RAW.mkdir(parents=True, exist_ok=True)
    run_dir = RUNS[condition]["run"]
    output = raw_path(condition, frame, slug)
    if output.exists():
        return output, "skipped existing output", 0

    before = {path.name for path in RAW.glob("image*.png")}
    cmd = [
        str(PLAY),
        "on",
        "image",
        f"frame={frame}",
        f"size={RENDER_SIZE},{RENDER_SIZE}",
        f"samples={SAMPLES}",
        "auto_scale=0",
        f"view_scale={CAMERA.view_scale}",
        f"zoom={CAMERA.zoom}",
        f"rotation={CAMERA.rotation}",
        "focus=0 0 0",
        "draw_memo=0",
        "label=none",
        "couple_select=0",
        "single_select=0",
        "draw_links=0",
        "back_color=white",
        "front_color=black",
        f"image_dir={RAW}",
        f"image_name={condition}_{CAMERA.key}_{slug}_frame{frame:04d}_raw",
        "image_format=png",
    ]
    proc = subprocess.run(cmd, cwd=run_dir, text=True, stdout=subprocess.PIPE, stderr=subprocess.STDOUT)
    if not output.exists():
        new_images = [path for path in RAW.glob("image*.png") if path.name not in before]
        if len(new_images) == 1:
            shutil.move(new_images[0], output)
    if not output.exists():
        raise RuntimeError(f"play did not produce {output}\nreturn={proc.returncode}\n{proc.stdout}")
    return output, " ".join(cmd), proc.returncode


def signal_from_raw(path: Path) -> np.ndarray:
    arr = np.asarray(Image.open(path).convert("RGB"), dtype=float) / 255.0
    lum = 0.299 * arr[..., 0] + 0.587 * arr[..., 1] + 0.114 * arr[..., 2]
    signal = np.clip(1.0 - lum, 0.0, 1.0)
    edge = max(35, signal.shape[0] // 42)
    signal[:edge, :] = 0.0
    signal[-edge:, :] = 0.0
    return signal


def union_crop(signals: list[np.ndarray], *, margin_fraction: float = 0.10) -> tuple[slice, slice]:
    mask = np.zeros_like(signals[0], dtype=bool)
    for signal in signals:
        mask |= signal > 0.018
    ys, xs = np.where(mask)
    if ys.size == 0 or xs.size == 0:
        return slice(0, signals[0].shape[0]), slice(0, signals[0].shape[1])
    height, width = signals[0].shape
    y0, y1 = int(ys.min()), int(ys.max()) + 1
    x0, x1 = int(xs.min()), int(xs.max()) + 1
    pad_y = int(round((y1 - y0) * margin_fraction))
    pad_x = int(round((x1 - x0) * margin_fraction))
    return slice(max(0, y0 - pad_y), min(height, y1 + pad_y)), slice(max(0, x0 - pad_x), min(width, x1 + pad_x))


def fluorescence_transform(signal: np.ndarray) -> np.ndarray:
    fine = gaussian_filter(signal, sigma=1.15)
    glow = gaussian_filter(signal, sigma=7.0)
    image = fine + 0.48 * glow
    positive = image[image > 0]
    scale = float(np.quantile(positive, 0.996)) if positive.size else 1.0
    image = np.power(np.clip(image / max(scale, 1e-12), 0.0, 1.0), 0.68)
    return np.dstack([0.025 * image, 0.92 * image, image])


def orient_image(rgb: np.ndarray) -> np.ndarray:
    return np.rot90(rgb, k=CAMERA.postprocess_rot90)


def render_records() -> list[dict[str, object]]:
    frame_times = list(PANEL_TIMEPOINTS)
    for idx, post_min in enumerate(np.linspace(0.0, FINAL_POST_MIN, SMOOTH_FRAME_COUNT)):
        frame_times.append((f"seq{idx:03d}", float(post_min), f"{post_min:.1f} min"))

    unique: dict[tuple[str, int, str], tuple[str, float, str]] = {}
    for slug, post_min, label in frame_times:
        frame = post_min_to_frame(post_min)
        for condition in RUNS:
            unique[(condition, frame, slug)] = (slug, post_min, label)

    raw_records: list[dict[str, object]] = []
    for (condition, frame, slug), (_slug, post_min, label) in sorted(unique.items()):
        raw, command, returncode = render_raw(condition, frame, slug)
        raw_records.append(
            {
                "condition": condition,
                "frame": frame,
                "slug": slug,
                "post_min": post_min,
                "label": label,
                "raw_path": raw,
                "command": command,
                "returncode": returncode,
            }
        )
    return raw_records


def process_images(raw_records: list[dict[str, object]]) -> None:
    RAW_ORIENTED.mkdir(parents=True, exist_ok=True)
    PROCESSED.mkdir(parents=True, exist_ok=True)
    signals = [signal_from_raw(Path(record["raw_path"])) for record in raw_records]
    crop_y, crop_x = union_crop(signals)
    for record, signal in zip(raw_records, signals):
        raw_rgb = np.asarray(Image.open(record["raw_path"]).convert("RGB"), dtype=float) / 255.0
        raw_cropped = raw_rgb[crop_y, crop_x]
        raw_oriented = orient_image(raw_cropped)
        raw_output = raw_oriented_path(str(record["condition"]), int(record["frame"]), str(record["slug"]))
        Image.fromarray(np.uint8(np.clip(raw_oriented, 0.0, 1.0) * 255)).save(raw_output)
        record["raw_oriented_path"] = raw_output

        cropped = signal[crop_y, crop_x]
        rgb = orient_image(fluorescence_transform(cropped))
        output = processed_path(str(record["condition"]), int(record["frame"]), str(record["slug"]))
        Image.fromarray(np.uint8(np.clip(rgb, 0.0, 1.0) * 255)).save(output)
        record["processed_path"] = output


def image_for(condition: str, slug: str, frame: int) -> Path:
    return processed_path(condition, frame, slug)


def raw_image_for(condition: str, slug: str, frame: int) -> Path:
    return raw_oriented_path(condition, frame, slug)


def make_snapshot_panel() -> Path:
    PANELS.mkdir(parents=True, exist_ok=True)
    fig, axes = plt.subplots(2, len(PANEL_TIMEPOINTS), figsize=(15.8, 7.8), constrained_layout=False)
    fig.patch.set_facecolor("black")
    for row, condition in enumerate(("nomotor_xlink", "rotatable_xlink")):
        for col, (slug, post_min, label) in enumerate(PANEL_TIMEPOINTS):
            frame = post_min_to_frame(post_min)
            ax = axes[row, col]
            ax.imshow(mpimg.imread(image_for(condition, slug, frame)))
            ax.set_axis_off()
            ax.set_facecolor("black")
            if row == 0:
                ax.text(
                    0.5,
                    1.035,
                    label,
                    transform=ax.transAxes,
                    ha="center",
                    va="bottom",
                    color="white",
                    fontsize=13,
                    weight="bold",
                )
            if col == 0:
                ax.text(
                    0.03,
                    0.96,
                    RUNS[condition]["label"],
                    transform=ax.transAxes,
                    ha="left",
                    va="top",
                    color="white",
                    fontsize=12,
                    weight="bold",
                    bbox={"boxstyle": "round,pad=0.20", "facecolor": "black", "edgecolor": "none", "alpha": 0.70},
                )
    fig.text(
        0.5,
        0.992,
        "Aggressive bottom depolymerization, continuous top actin supply",
        ha="center",
        va="top",
        color="white",
        fontsize=14,
    )
    fig.text(0.5, 0.955, "time after analysis onset", ha="center", va="top", color="#cfcfcf", fontsize=10)
    fig.subplots_adjust(left=0.012, right=0.988, top=0.89, bottom=0.025, wspace=0.02, hspace=0.07)
    output = PANELS / "aggressive_turnover_no_motor_vs_motor_snapshots.png"
    fig.savefig(output, dpi=500, bbox_inches="tight", facecolor="black")
    fig.savefig(output.with_suffix(".pdf"), bbox_inches="tight", facecolor="black")
    plt.close(fig)
    return output


def label_image(path: Path, title: str, time_label: str, target_w: int = MOVIE_PANEL_W, target_h: int = MOVIE_PANEL_H) -> Image.Image:
    img = Image.open(path).convert("RGB")
    background = tuple(int(v) for v in ImageStat.Stat(img).mean)
    background = (255, 255, 255) if sum(background) > 480 else (0, 0, 0)
    scale = min(target_w / img.width, target_h / img.height)
    resized = img.resize((int(round(img.width * scale)), int(round(img.height * scale))), Image.Resampling.LANCZOS)
    content = Image.new("RGB", (target_w, target_h), background)
    content.paste(resized, ((target_w - resized.width) // 2, (target_h - resized.height) // 2))
    canvas = Image.new("RGB", (target_w, target_h + 82), "black")
    canvas.paste(content, (0, 82))
    draw = ImageDraw.Draw(canvas)
    try:
        font_big = ImageFont.truetype("DejaVuSans-Bold.ttf", 34)
        font_small = ImageFont.truetype("DejaVuSans.ttf", 24)
    except OSError:
        font_big = ImageFont.load_default()
        font_small = ImageFont.load_default()
    draw.text((20, 10), title, fill="white", font=font_big)
    draw.text((20, 50), time_label, fill=(190, 190, 190), font=font_small)
    return canvas


def save_gif(frames: list[Image.Image], output: Path, duration_ms: int = 120) -> None:
    output.parent.mkdir(parents=True, exist_ok=True)
    frames[0].save(
        output,
        save_all=True,
        append_images=frames[1:],
        duration=duration_ms,
        loop=0,
        optimize=False,
    )


def make_frame_sequences() -> dict[str, Path]:
    FRAMES.mkdir(parents=True, exist_ok=True)
    outputs: dict[str, Path] = {}

    for style, lookup in (("cyan", image_for), ("raw", raw_image_for)):
        per_condition_frames: dict[str, list[Image.Image]] = {condition: [] for condition in RUNS}
        side_by_side: list[Image.Image] = []

        for idx, post_min in enumerate(np.linspace(0.0, FINAL_POST_MIN, SMOOTH_FRAME_COUNT)):
            slug = f"seq{idx:03d}"
            frame = post_min_to_frame(float(post_min))
            time_label = f"{post_min:.1f} min after onset"
            labeled = {}
            for condition, spec in RUNS.items():
                img = label_image(lookup(condition, slug, frame), str(spec["label"]), time_label)
                labeled[condition] = img
                per_condition_frames[condition].append(img)
                img.save(FRAMES / f"{idx:03d}_{condition}_{style}_{slug}.png")

            gap = 24
            left = labeled["nomotor_xlink"]
            right = labeled["rotatable_xlink"]
            height = max(left.height, right.height)
            width = left.width + right.width + gap
            canvas = Image.new("RGB", (width, height), "black")
            canvas.paste(left, (0, 0))
            canvas.paste(right, (left.width + gap, 0))
            side_by_side.append(canvas)
            canvas.save(FRAMES / f"{idx:03d}_side_by_side_{style}_{slug}.png")

        for condition in RUNS:
            output = FRAMES / f"{condition}_aggressive_turnover_{style}.gif"
            save_gif(per_condition_frames[condition], output)
            outputs[f"{condition}_{style}"] = output
        combined = FRAMES / f"aggressive_turnover_no_motor_vs_motor_side_by_side_{style}.gif"
        save_gif(side_by_side, combined)
        outputs[f"side_by_side_{style}"] = combined

    legacy_combined = FRAMES / "aggressive_turnover_no_motor_vs_motor_side_by_side.gif"
    shutil.copyfile(outputs["side_by_side_cyan"], legacy_combined)
    outputs["side_by_side"] = legacy_combined
    for condition in RUNS:
        legacy = FRAMES / f"{condition}_aggressive_turnover.gif"
        shutil.copyfile(outputs[f"{condition}_cyan"], legacy)
        outputs[condition] = legacy
    return outputs


def cumulative_auc(time_min: np.ndarray, speed_um_s: np.ndarray) -> np.ndarray:
    auc = np.zeros_like(time_min, dtype=float)
    for i in range(1, time_min.size):
        dt_s = (time_min[i] - time_min[i - 1]) * 60.0
        auc[i] = auc[i - 1] + 0.5 * (speed_um_s[i] + speed_um_s[i - 1]) * dt_s
    return auc


def load_condition_timecourses() -> list[dict[str, object]]:
    rows = []
    for row in read_csv(ANALYSIS / "condition_timecourses.csv"):
        if row["scenario"] != "large_long" or row["condition"] not in RUNS:
            continue
        post_min = f(row["time_min"]) - ONSET_MIN
        if post_min < -1e-9:
            continue
        item = dict(row)
        item["post_onset_min"] = post_min
        rows.append(item)
    return rows


def group_condition_timecourses(rows: list[dict[str, object]]) -> dict[str, list[dict[str, object]]]:
    grouped: dict[str, list[dict[str, object]]] = {}
    for row in rows:
        grouped.setdefault(str(row["condition"]), []).append(row)
    for condition in grouped:
        grouped[condition].sort(key=lambda row: f(row["post_onset_min"]))
    return grouped


def add_phase_spans(ax: plt.Axes) -> None:
    colors = ["#f0f0f0", "#ffffff", "#f0f0f0"]
    for idx, (_key, label, lo, hi) in enumerate(PHASES):
        ax.axvspan(lo, hi, color=colors[idx], zorder=-2)
        ax.text((lo + hi) / 2, 1.015, label, transform=ax.get_xaxis_transform(), ha="center", va="bottom", fontsize=9, color="#555555")


def plot_timecourses() -> Path:
    PLOTS.mkdir(parents=True, exist_ok=True)
    rows = load_condition_timecourses()
    grouped = group_condition_timecourses(rows)
    fig, axes = plt.subplots(2, 2, figsize=(11.6, 7.6), sharex=True)
    specs = [
        ("vz_abs", "Mean |v_z| (um/s)", False),
        ("vz_abs", "Cumulative AUC |v_z| (um)", True),
        ("chewer_band_mass_fraction", "Chewer-band mass fraction", False),
        ("bottom_half_mass_fraction", "Bottom-half mass fraction", False),
    ]
    for ax, (field, ylabel, is_auc), label in zip(axes.flat, specs, ["A", "B", "C", "D"]):
        add_phase_spans(ax)
        for condition, spec in RUNS.items():
            cr = grouped[condition]
            t = np.asarray([f(row["post_onset_min"]) for row in cr], dtype=float)
            mean = np.asarray([f(row[f"{field}_mean"]) for row in cr], dtype=float)
            sem_values = np.asarray([f(row.get(f"{field}_sem"), 0.0) for row in cr], dtype=float)
            if is_auc:
                mean = cumulative_auc(t, mean)
                sem_values = np.zeros_like(mean)
            ax.plot(t, mean, color=str(spec["color"]), linewidth=2.4, label=str(spec["label"]))
            if not is_auc:
                ax.fill_between(t, mean - sem_values, mean + sem_values, color=str(spec["color"]), alpha=0.18, linewidth=0)
        ax.text(0.012, 0.98, label, transform=ax.transAxes, ha="left", va="top", fontsize=15, fontweight="bold")
        ax.set_ylabel(ylabel)
        ax.grid(axis="y", color="#d8d8d8", linewidth=0.8)
        ax.set_xlim(0, FINAL_POST_MIN)
    for ax in axes[-1, :]:
        ax.set_xlabel("Time after onset (min)")
    axes[0, 1].legend(loc="upper left", bbox_to_anchor=(1.02, 1.0), frameon=False)
    fig.tight_layout(w_pad=2.2, h_pad=2.0)
    output = PLOTS / "aggressive_post_onset_timecourses.png"
    fig.savefig(output, dpi=300, bbox_inches="tight")
    fig.savefig(output.with_suffix(".svg"), bbox_inches="tight")
    plt.close(fig)
    return output


def integrate_phase(rows: list[dict[str, str]], metric: str) -> float:
    sorted_rows = sorted(rows, key=lambda row: f(row["time_min"]))
    t = np.asarray([f(row["time_min"]) - ONSET_MIN for row in sorted_rows], dtype=float)
    y = np.asarray([f(row[metric]) for row in sorted_rows], dtype=float)
    if t.size < 2:
        return float("nan")
    return float(np.trapezoid(y, t * 60.0))


def make_phase_summary() -> tuple[Path, Path]:
    run_rows = [
        row
        for row in read_csv(ANALYSIS / "run_timecourses_long.csv")
        if row["scenario"] == "large_long" and row["condition"] in RUNS and f(row["time_min"]) >= ONSET_MIN
    ]

    per_run_phase: list[dict[str, object]] = []
    for condition in RUNS:
        run_ids = sorted({row["run_dir"] for row in run_rows if row["condition"] == condition})
        for run_id in run_ids:
            rr = [row for row in run_rows if row["condition"] == condition and row["run_dir"] == run_id]
            for phase_key, phase_label, lo, hi in PHASES:
                pr = [row for row in rr if lo <= f(row["time_min"]) - ONSET_MIN <= hi]
                if not pr:
                    continue
                per_run_phase.append(
                    {
                        "condition": condition,
                        "condition_label": RUNS[condition]["label"],
                        "run_dir": run_id,
                        "phase": phase_key,
                        "phase_label": phase_label,
                        "mean_vz_abs": float(np.nanmean([f(row["vz_abs"]) for row in pr])),
                        "auc_vz_abs": integrate_phase(pr, "vz_abs"),
                        "mean_chewer_band_mass_fraction": float(np.nanmean([f(row["chewer_band_mass_fraction"]) for row in pr])),
                        "mean_bottom_half_mass_fraction": float(np.nanmean([f(row["bottom_half_mass_fraction"]) for row in pr])),
                        "mean_speckle_length_proxy_um": float(np.nanmean([f(row["speckle_length_proxy_um"]) for row in pr])),
                    }
                )

    summary: list[dict[str, object]] = []
    for condition in RUNS:
        for phase_key, phase_label, _lo, _hi in PHASES:
            rows = [row for row in per_run_phase if row["condition"] == condition and row["phase"] == phase_key]
            if not rows:
                continue
            item: dict[str, object] = {
                "condition": condition,
                "condition_label": RUNS[condition]["label"],
                "phase": phase_key,
                "phase_label": phase_label,
                "n": len(rows),
            }
            for metric in (
                "mean_vz_abs",
                "auc_vz_abs",
                "mean_chewer_band_mass_fraction",
                "mean_bottom_half_mass_fraction",
                "mean_speckle_length_proxy_um",
            ):
                values = [f(row[metric]) for row in rows]
                item[f"{metric}_mean"] = float(np.nanmean(values))
                item[f"{metric}_sem"] = sem(values)
            summary.append(item)

    write_csv(OUT / "phase_metrics_per_run.csv", per_run_phase)
    write_csv(OUT / "phase_metrics_summary.csv", summary)

    fig, axes = plt.subplots(2, 2, figsize=(10.8, 7.6))
    metrics = [
        ("mean_vz_abs", "Mean |v_z| (um/s)"),
        ("auc_vz_abs", "AUC |v_z| (um)"),
        ("mean_chewer_band_mass_fraction", "Mean chewer-band fraction"),
        ("mean_bottom_half_mass_fraction", "Mean bottom-half fraction"),
    ]
    phase_labels = [phase[1] for phase in PHASES]
    x = np.arange(len(PHASES))
    width = 0.34
    for ax, (metric, ylabel), panel in zip(axes.flat, metrics, ["A", "B", "C", "D"]):
        for offset, condition in [(-width / 2, "nomotor_xlink"), (width / 2, "rotatable_xlink")]:
            values = []
            errors = []
            for phase_key, _phase_label, _lo, _hi in PHASES:
                row = next(row for row in summary if row["condition"] == condition and row["phase"] == phase_key)
                values.append(f(row[f"{metric}_mean"]))
                errors.append(f(row[f"{metric}_sem"], 0.0))
            ax.bar(
                x + offset,
                values,
                width,
                yerr=errors,
                color=str(RUNS[condition]["color"]),
                edgecolor="#222222",
                linewidth=0.8,
                capsize=3,
                label=str(RUNS[condition]["label"]),
            )
        ax.text(0.012, 0.98, panel, transform=ax.transAxes, ha="left", va="top", fontsize=15, fontweight="bold")
        ax.set_ylabel(ylabel)
        ax.set_xticks(x, phase_labels)
        ax.grid(axis="y", color="#d8d8d8", linewidth=0.8)
        ax.set_axisbelow(True)
    axes[0, 1].legend(loc="upper left", bbox_to_anchor=(1.02, 1.0), frameon=False)
    fig.tight_layout(w_pad=2.2, h_pad=2.0)
    output = PLOTS / "aggressive_phase_summary.png"
    fig.savefig(output, dpi=300, bbox_inches="tight")
    fig.savefig(output.with_suffix(".svg"), bbox_inches="tight")
    plt.close(fig)
    return OUT / "phase_metrics_summary.csv", output


COLORS = {
    "ink": RGBColor(22, 24, 29),
    "muted": RGBColor(82, 89, 99),
    "teal": RGBColor(10, 159, 122),
    "gray": RGBColor(232, 235, 238),
    "light": RGBColor(247, 248, 250),
}
WIDE = Inches(13.333)
HIGH = Inches(7.5)


def add_textbox(slide, text: str, left, top, width, height, *, size=22, bold=False, color="ink", align=None):
    shape = slide.shapes.add_textbox(left, top, width, height)
    frame = shape.text_frame
    frame.clear()
    p = frame.paragraphs[0]
    p.text = text
    p.font.size = Pt(size)
    p.font.bold = bold
    p.font.color.rgb = COLORS[color] if isinstance(color, str) else color
    p.font.name = "Aptos"
    if align is not None:
        p.alignment = align
    return shape


def add_title(slide, title: str, subtitle: str | None = None) -> None:
    add_textbox(slide, title, Inches(0.55), Inches(0.25), Inches(12.2), Inches(0.48), size=25, bold=True)
    if subtitle:
        add_textbox(slide, subtitle, Inches(0.58), Inches(0.76), Inches(12.0), Inches(0.3), size=12, color="muted")
    line = slide.shapes.add_shape(1, Inches(0.55), Inches(1.08), Inches(12.2), Inches(0.02))
    line.fill.solid()
    line.fill.fore_color.rgb = COLORS["gray"]
    line.line.fill.background()


def add_bullets(slide, bullets: list[str], left, top, width, height, *, size=19) -> None:
    shape = slide.shapes.add_textbox(left, top, width, height)
    frame = shape.text_frame
    frame.clear()
    frame.margin_left = Inches(0.03)
    frame.margin_right = Inches(0.03)
    for idx, bullet in enumerate(bullets):
        p = frame.paragraphs[0] if idx == 0 else frame.add_paragraph()
        p.text = bullet if bullet.startswith("- ") else f"- {bullet}"
        p.level = 0
        p.font.size = Pt(size)
        p.font.name = "Aptos"
        p.font.color.rgb = COLORS["ink"]
        p.space_after = Pt(7)


def image_fit(path: Path, box_left, box_top, box_width, box_height):
    with Image.open(path) as img:
        iw, ih = img.size
    aspect = iw / ih
    box_aspect = box_width / box_height
    if aspect > box_aspect:
        width = box_width
        height = box_width / aspect
    else:
        height = box_height
        width = box_height * aspect
    left = box_left + (box_width - width) / 2
    top = box_top + (box_height - height) / 2
    return left, top, width, height


def add_image(slide, path: Path, left, top, width, height) -> None:
    l, t, w, h = image_fit(path, left, top, width, height)
    slide.shapes.add_picture(str(path), l, t, width=w, height=h)


def blank(prs: Presentation):
    return prs.slides.add_slide(prs.slide_layouts[6])


def prepare_reference_assets() -> dict[str, Path]:
    ASSETS.mkdir(parents=True, exist_ok=True)
    assets: dict[str, Path] = {}
    if REFERENCE_BACKGROUND_SCREENSHOT.exists():
        with Image.open(REFERENCE_BACKGROUND_SCREENSHOT).convert("RGB") as img:
            width, height = img.size
            crop = img.crop(
                (
                    int(width * 0.53),
                    int(height * 0.14),
                    int(width * 0.735),
                    int(height * 0.895),
                )
            )
            output = ASSETS / "experimental_reference_crop.png"
            crop.save(output)
            assets["experimental_reference"] = output
    return assets


def make_methods_schematic() -> Path:
    ASSETS.mkdir(parents=True, exist_ok=True)
    method_reference = raw_image_for("rotatable_xlink", "t00", post_min_to_frame(0.0))
    fig = plt.figure(figsize=(8.0, 5.0), facecolor="white")

    side = fig.add_axes([0.04, 0.08, 0.43, 0.84])
    side.imshow(mpimg.imread(method_reference))
    side.set_axis_off()
    side.add_patch(patches.Rectangle((0.08, 0.69), 0.84, 0.10, transform=side.transAxes, facecolor="#d9f2ff", edgecolor="none", alpha=0.38))
    side.add_patch(patches.Rectangle((0.08, 0.10), 0.84, 0.28, transform=side.transAxes, facecolor="#ffb3a7", edgecolor="none", alpha=0.36))
    side.annotate(
        "aligned actin\npulses every 1 min",
        xy=(0.50, 0.82),
        xytext=(0.50, 0.97),
        xycoords="axes fraction",
        textcoords="axes fraction",
        ha="center",
        va="top",
        fontsize=10,
        color="#075985",
        arrowprops={"arrowstyle": "->", "lw": 1.8, "color": "#075985"},
    )
    side.annotate(
        "bottom minus-end\nchewing zone",
        xy=(0.56, 0.22),
        xytext=(0.98, 0.22),
        xycoords="axes fraction",
        textcoords="axes fraction",
        ha="left",
        va="center",
        fontsize=10,
        color="#9a3412",
        arrowprops={"arrowstyle": "->", "lw": 1.8, "color": "#9a3412"},
    )
    side.annotate(
        "",
        xy=(0.03, 0.12),
        xytext=(0.03, 0.88),
        xycoords="axes fraction",
        arrowprops={"arrowstyle": "<->", "lw": 1.7, "color": "#0f5f8a"},
    )
    side.text(0.00, 0.50, "30 um", transform=side.transAxes, ha="right", va="center", fontsize=10, color="#0f5f8a")

    top = fig.add_axes([0.55, 0.25, 0.34, 0.50])
    top.set_aspect("equal")
    top.set_xlim(-1.15, 1.15)
    top.set_ylim(-1.15, 1.15)
    top.set_axis_off()
    top.add_patch(patches.Circle((0, 0), 1.0, facecolor="none", edgecolor="#0f5f8a", lw=2.0))
    top.add_patch(patches.Circle((0, 0), 0.91, facecolor="none", edgecolor="#0f5f8a", lw=2.0))
    top.annotate("", xy=(0.91, 0.15), xytext=(1.0, 0.28), arrowprops={"arrowstyle": "<->", "lw": 1.5, "color": "#0f5f8a"})
    top.text(0.94, 0.38, "0.5 um shell", ha="center", va="bottom", fontsize=10, color="#0f5f8a")
    top.annotate("", xy=(0, -0.91), xytext=(0, 0.91), arrowprops={"arrowstyle": "<->", "lw": 1.6, "color": "#0f5f8a"})
    top.text(0.08, 0.0, "inner diameter\n10.5 um", ha="left", va="center", fontsize=10, color="#222222")
    top.text(0, -1.18, "annular shell cross-section", ha="center", va="top", fontsize=10, color="#555555")

    fig.text(0.04, 0.97, "Scaled turnover test geometry", ha="left", va="top", fontsize=13, weight="bold")
    fig.text(0.55, 0.90, "Source-sink design", ha="left", va="top", fontsize=13, weight="bold")
    output = ASSETS / "turnover_methods_schematic.png"
    fig.savefig(output, dpi=300, bbox_inches="tight", facecolor="white")
    plt.close(fig)
    return output


def build_deck(
    snapshot_panel: Path,
    gifs: dict[str, Path],
    timecourse_plot: Path,
    phase_plot: Path,
    reference_assets: dict[str, Path],
    methods_schematic: Path,
) -> Path:
    prs = Presentation()
    prs.slide_width = WIDE
    prs.slide_height = HIGH
    cyan_gif = gifs["side_by_side_cyan"]
    raw_gif = gifs["side_by_side_raw"]
    sim_reference = image_for("rotatable_xlink", "t02", post_min_to_frame(2.0))

    s = blank(prs)
    add_textbox(s, "Continuous actin supply with aggressive bottom turnover", Inches(0.65), Inches(0.55), Inches(12.0), Inches(0.8), size=34, bold=True)
    add_textbox(s, "Large scaled annulus, no-motor versus wall-bound motor clusters", Inches(0.70), Inches(1.28), Inches(11.8), Inches(0.38), size=17, color="muted")
    add_bullets(
        s,
        [
            "The setup removes the pre-onset equilibration phase from result metrics.",
            "Top actin supply is pulsed every 1 min; bottom minus-end chewing is 2x chewer count and 2x chewing speed.",
            "Only the aggressive turnover case is shown because it gives the clearest bottom depolymerization phenotype.",
        ],
        Inches(0.80),
        Inches(2.05),
        Inches(11.8),
        Inches(3.2),
        size=21,
    )

    s = blank(prs)
    add_textbox(s, "Background", Inches(0.70), Inches(0.50), Inches(7.2), Inches(0.7), size=40, bold=True)
    add_bullets(
        s,
        [
            "Membrane-bound myosin clusters are observed near actin structures.",
            "Earlier simulations with random initial filament orientations did not produce strong directional actin motion.",
            "This test asks whether continuous top actin supply plus bottom depolymerization changes the motor-driven behavior.",
        ],
        Inches(0.72),
        Inches(2.05),
        Inches(5.9),
        Inches(4.2),
        size=24,
    )
    if "experimental_reference" in reference_assets:
        add_image(s, reference_assets["experimental_reference"], Inches(7.10), Inches(1.28), Inches(2.75), Inches(5.45))
        add_textbox(s, "Experimental reference", Inches(7.10), Inches(6.88), Inches(2.75), Inches(0.24), size=10, color="muted", align=PP_ALIGN.CENTER)
    add_image(s, sim_reference, Inches(10.25), Inches(1.28), Inches(2.55), Inches(5.45))
    add_textbox(s, "Simulation rendering", Inches(10.25), Inches(6.88), Inches(2.55), Inches(0.24), size=10, color="muted", align=PP_ALIGN.CENTER)

    s = blank(prs)
    add_textbox(s, "Simulation Design and Methods", Inches(0.70), Inches(0.42), Inches(11.9), Inches(0.7), size=38, bold=True)
    add_image(s, methods_schematic, Inches(0.45), Inches(1.35), Inches(5.70), Inches(5.55))
    add_bullets(
        s,
        [
            "Smaller long annular shell used for fast testing: height 30 um, inner radius 5.25 um, outer radius 5.75 um.",
            "This is a source-sink approximation of biological turnover: aligned actin is supplied near the top every 1 min.",
            "Bottom 30% of the annulus contains aggressive minus-end chewers: 2x chewer count and 2x chewing speed.",
            "Crosslinkers are replenished with each supply pulse so new actin can join the existing network.",
            "Main comparison: same turnover design with no motors versus rotatable membrane-bound motor clusters.",
            "Analysis starts after the initial 1 min setup period so motor-introduction transients do not dominate the plots.",
        ],
        Inches(6.35),
        Inches(1.42),
        Inches(6.55),
        Inches(4.95),
        size=18,
    )

    s = blank(prs)
    add_title(s, "Representative Post-Onset Snapshots", "Median-AUC replicates; cyan fluorescence rendering of Cytosim actin")
    add_image(s, snapshot_panel, Inches(0.30), Inches(1.25), Inches(12.75), Inches(5.75))

    s = blank(prs)
    add_title(s, "Processed Side-By-Side Movie", f"Animated GIF generated from {SMOOTH_FRAME_COUNT} post-onset frames")
    add_image(s, cyan_gif, Inches(0.60), Inches(1.18), Inches(12.05), Inches(5.9))
    add_textbox(s, f"GIF asset: {cyan_gif.relative_to(ROOT)}", Inches(0.70), Inches(6.92), Inches(11.8), Inches(0.28), size=10, color="muted", align=PP_ALIGN.CENTER)

    s = blank(prs)
    add_title(s, "Raw Cytosim Movie Frames", f"Unprocessed Cytosim line rendering from the same {SMOOTH_FRAME_COUNT} sampled frames")
    add_image(s, raw_gif, Inches(0.60), Inches(1.18), Inches(12.05), Inches(5.9))
    add_textbox(s, f"GIF asset: {raw_gif.relative_to(ROOT)}", Inches(0.70), Inches(6.92), Inches(11.8), Inches(0.28), size=10, color="muted", align=PP_ALIGN.CENTER)

    s = blank(prs)
    add_title(s, "Post-Onset Timecourses", "Phase shading: early, middle, late; all pre-onset points removed")
    add_image(s, timecourse_plot, Inches(0.45), Inches(1.18), Inches(12.25), Inches(5.85))

    s = blank(prs)
    add_title(s, "Phase Summary", "Each bar summarizes replicate-level post-onset trajectories")
    add_image(s, phase_plot, Inches(0.45), Inches(1.18), Inches(12.25), Inches(5.85))

    s = blank(prs)
    add_title(s, "What The Timecourses Show", "Readout: post-onset dynamics after the initial setup period")
    add_bullets(
        s,
        [
            "The motor condition sustains higher axial actin motion than the no-motor condition after the source-sink turnover has started.",
            "The cumulative |v_z| curve is useful here because the phenotype is dynamic: small persistent velocity differences add up across the run.",
            "Mass-fraction traces show that aggressive bottom turnover creates a real sink, while top pulses keep supplying actin instead of letting the system simply deplete.",
        ],
        Inches(0.90),
        Inches(1.55),
        Inches(11.75),
        Inches(4.9),
        size=24,
    )

    s = blank(prs)
    add_title(s, "What The Phase Summary Adds", "Readout: early, middle, and late behavior summarized per replicate")
    add_bullets(
        s,
        [
            "Splitting the run into phases separates the early reorganization from the later quasi-steady behavior.",
            "The motor case is expected to matter most when the actin network is still connected enough to transmit force through crosslinkers.",
            "The no-motor case is the turnover-only baseline: it tells us what bottom chewing plus top supply can do without active wall-bound motors.",
        ],
        Inches(0.90),
        Inches(1.55),
        Inches(11.75),
        Inches(4.9),
        size=24,
    )

    s = blank(prs)
    add_title(s, "Interpretation For The Full-Size Run", "Why this scaled result is worth carrying forward")
    add_bullets(
        s,
        [
            "This scaled system is not the final biological comparison; it is a controlled test showing that continuous supply and bottom depolymerization can coexist with motor-driven flow.",
            "The next production run should restore the comparable central-cell scale and filament density while keeping the same source-sink logic.",
            "If the full-size system preserves the motor/no-motor separation, it becomes a cleaner model evolution from polarity-only simulations to turnover-driven central-cell dynamics.",
        ],
        Inches(0.90),
        Inches(1.55),
        Inches(11.75),
        Inches(4.9),
        size=23,
    )

    s = blank(prs)
    add_title(s, "Next Full-Size Submission", "Use this result to motivate a comparable-scale production run")
    add_bullets(
        s,
        [
            "Keep the same turnover logic: top supply every 1 min and aggressive bottom minus-end chewing.",
            "Restore the older comparable annulus scale and filament density so it can sit beside the earlier central-cell simulations.",
            "Submit no-motor and motor cases first; expand only after the full-size behavior is visually and quantitatively stable.",
        ],
        Inches(0.85),
        Inches(1.55),
        Inches(11.7),
        Inches(4.6),
        size=23,
    )

    output = OUT / "aggressive_turnover_large_results.pptx"
    prs.save(output)
    return output


def write_readme(
    raw_records: list[dict[str, object]],
    snapshot_panel: Path,
    methods_schematic: Path,
    gifs: dict[str, Path],
    timecourse_plot: Path,
    phase_plot: Path,
    deck: Path,
) -> None:
    write_csv(
        OUT / "render_log.csv",
        [
            {
                "condition": record["condition"],
                "slug": record["slug"],
                "frame": record["frame"],
                "post_min": record["post_min"],
                "returncode": record["returncode"],
                "raw_path": Path(record["raw_path"]).relative_to(ROOT).as_posix(),
                "raw_oriented_path": Path(record["raw_oriented_path"]).relative_to(ROOT).as_posix(),
                "processed_path": Path(record["processed_path"]).relative_to(ROOT).as_posix(),
            }
            for record in raw_records
        ],
    )
    lines = [
        "# Aggressive Turnover Large Visual Deck",
        "",
        "Scope:",
        "- Source run set: `turnover_continuous_supply_scaled/job12`.",
        "- Variant: `chewer_count2x_rate2x`.",
        "- Included only `large_long` no-motor and motor conditions.",
        "- Pre-onset time points are omitted from the result plots and phase metrics.",
        "",
        "Representative runs:",
        "- No motors: `job12/save/r0003`, median post-onset AUC |v_z| replicate.",
        "- Motors: `job12/save/r0007`, median post-onset AUC |v_z| replicate.",
        "",
        "Main assets:",
        f"- Methods schematic: `{methods_schematic.relative_to(ROOT)}`",
        f"- Snapshot panel: `{snapshot_panel.relative_to(ROOT)}`",
        f"- Processed side-by-side GIF: `{gifs['side_by_side_cyan'].relative_to(ROOT)}`",
        f"- Raw Cytosim side-by-side GIF: `{gifs['side_by_side_raw'].relative_to(ROOT)}`",
        f"- Timecourse plot: `{timecourse_plot.relative_to(ROOT)}`",
        f"- Phase plot: `{phase_plot.relative_to(ROOT)}`",
        f"- Slides: `{deck.relative_to(ROOT)}`",
        f"- Render log: `{(OUT / 'render_log.csv').relative_to(ROOT)}`",
        "",
        "Rendered frames:",
    ]
    for key, path in gifs.items():
        lines.append(f"- `{key}`: `{path.relative_to(ROOT)}`")
    lines.extend(
        [
            "",
            "Render note:",
            f"- Movies are sampled with `{SMOOTH_FRAME_COUNT}` post-onset frames from 0 to {FINAL_POST_MIN:g} min after onset.",
            "- The reused `job12/save/play` binary may return status 11 after writing a PNG on this machine.",
            "- All listed raw and processed images were created successfully; the compact `render_log.csv` records frame numbers and return codes.",
        ]
    )
    (OUT / "README.md").write_text("\n".join(lines) + "\n", encoding="utf-8")


def main() -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    reset_generated_outputs()
    raw_records = render_records()
    process_images(raw_records)
    reference_assets = prepare_reference_assets()
    methods_schematic = make_methods_schematic()
    snapshot_panel = make_snapshot_panel()
    gifs = make_frame_sequences()
    timecourse_plot = plot_timecourses()
    _summary_csv, phase_plot = make_phase_summary()
    deck = build_deck(snapshot_panel, gifs, timecourse_plot, phase_plot, reference_assets, methods_schematic)
    write_readme(raw_records, snapshot_panel, methods_schematic, gifs, timecourse_plot, phase_plot, deck)
    print(OUT)


if __name__ == "__main__":
    main()
