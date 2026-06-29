#!/usr/bin/env python3
from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt
from PIL import Image


ROOT = Path(__file__).resolve().parent.parent
ASSET_DIR = ROOT / "analysis" / "results" / "clu_minusz_slides_2026-06-02"
OUT = ASSET_DIR / "clu_minusz_initial_alignment_slides.pptx"
NOTES = ASSET_DIR / "clu_minusz_initial_alignment_slide_notes.md"

WIDE = Inches(13.333)
HIGH = Inches(7.5)

COLORS = {
    "ink": RGBColor(22, 24, 29),
    "muted": RGBColor(82, 89, 99),
    "teal": RGBColor(27, 121, 103),
    "orange": RGBColor(217, 95, 2),
    "purple": RGBColor(117, 112, 179),
    "gray": RGBColor(232, 235, 238),
    "light": RGBColor(247, 248, 250),
}


def add_textbox(slide, text, left, top, width, height, *, size=22, bold=False, color="ink", align=None):
    shape = slide.shapes.add_textbox(left, top, width, height)
    frame = shape.text_frame
    frame.clear()
    p = frame.paragraphs[0]
    p.text = text
    p.font.size = Pt(size)
    p.font.bold = bold
    p.font.color.rgb = COLORS[color] if isinstance(color, str) else color
    p.font.name = "Aptos"
    if align is not None:
        p.alignment = align
    return shape


def add_title(slide, title, subtitle=None):
    add_textbox(slide, title, Inches(0.55), Inches(0.25), Inches(12.2), Inches(0.48), size=26, bold=True)
    if subtitle:
        add_textbox(slide, subtitle, Inches(0.58), Inches(0.75), Inches(12.0), Inches(0.30), size=12, color="muted")
    line = slide.shapes.add_shape(1, Inches(0.55), Inches(1.08), Inches(12.2), Inches(0.02))
    line.fill.solid()
    line.fill.fore_color.rgb = COLORS["gray"]
    line.line.fill.background()


def add_footer(slide, number):
    add_textbox(
        slide,
        f"Initially aligned minus-z simulations | {number}",
        Inches(0.55),
        Inches(7.12),
        Inches(12.2),
        Inches(0.22),
        size=8,
        color="muted",
        align=PP_ALIGN.RIGHT,
    )


def add_bullets(slide, bullets, left, top, width, height, *, size=20, color="ink", gap=0.16):
    shape = slide.shapes.add_textbox(left, top, width, height)
    frame = shape.text_frame
    frame.clear()
    frame.margin_left = Inches(0.03)
    frame.margin_right = Inches(0.03)
    frame.margin_top = Inches(0.01)
    frame.margin_bottom = Inches(0.01)
    for idx, bullet in enumerate(bullets):
        p = frame.paragraphs[0] if idx == 0 else frame.add_paragraph()
        p.text = bullet
        p.level = 0
        p.font.size = Pt(size)
        p.font.name = "Aptos"
        p.font.color.rgb = COLORS[color]
        p.space_after = Pt(gap * 72)
    return shape


def add_label(slide, text, left, top, width, *, color="teal"):
    shape = slide.shapes.add_shape(1, left, top, width, Inches(0.34))
    shape.fill.solid()
    shape.fill.fore_color.rgb = COLORS[color]
    shape.line.fill.background()
    tf = shape.text_frame
    tf.clear()
    p = tf.paragraphs[0]
    p.text = text
    p.alignment = PP_ALIGN.CENTER
    p.font.size = Pt(12)
    p.font.bold = True
    p.font.color.rgb = RGBColor(255, 255, 255)
    p.font.name = "Aptos"
    return shape


def image_fit(path: Path, box_left, box_top, box_width, box_height):
    with Image.open(path) as img:
        iw, ih = img.size
    aspect = iw / ih
    box_aspect = box_width / box_height
    if aspect > box_aspect:
        width = box_width
        height = box_width / aspect
    else:
        height = box_height
        width = box_height * aspect
    left = box_left + (box_width - width) / 2
    top = box_top + (box_height - height) / 2
    return left, top, width, height


def add_image(slide, rel_path, left, top, width, height):
    path = ASSET_DIR / rel_path
    l, t, w, h = image_fit(path, left, top, width, height)
    slide.shapes.add_picture(str(path), l, t, width=w, height=h)


def add_card(slide, title, body, left, top, width, height, *, accent="teal"):
    shape = slide.shapes.add_shape(1, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = COLORS["light"]
    shape.line.color.rgb = COLORS["gray"]
    shape.line.width = Pt(1)
    add_label(slide, title, left, top, width, color=accent)
    add_bullets(slide, body, left + Inches(0.18), top + Inches(0.52), width - Inches(0.36), height - Inches(0.64), size=14)


def blank(prs):
    return prs.slides.add_slide(prs.slide_layouts[6])


def build_deck() -> None:
    prs = Presentation()
    prs.slide_width = WIDE
    prs.slide_height = HIGH
    notes: list[tuple[str, list[str]]] = []

    slide_no = 1

    s = blank(prs)
    add_textbox(s, "Initially aligned actin filaments", Inches(0.65), Inches(0.55), Inches(12.0), Inches(0.7), size=38, bold=True)
    add_textbox(
        s,
        "Rotatable vs fixed-orientation wall-bound myosin clusters in an annular domain",
        Inches(0.70),
        Inches(1.28),
        Inches(11.8),
        Inches(0.38),
        size=18,
        color="muted",
    )
    add_card(
        s,
        "Question",
        [
            "If filament polarity is supplied at t = 0, how strongly do wall-bound motor clusters drive coherent axial motion?",
            "Does motor rotatability change the outcome when the actin starts aligned along minus-z?",
        ],
        Inches(0.75),
        Inches(2.15),
        Inches(5.7),
        Inches(2.35),
        accent="teal",
    )
    add_card(
        s,
        "Readout",
        [
            "Movies first, then unwrapped heatmaps to show spatial reorganization.",
            "Summary metrics quantify speed, persistence, and active transport across the sweep.",
        ],
        Inches(6.9),
        Inches(2.15),
        Inches(5.7),
        Inches(2.35),
        accent="orange",
    )
    add_textbox(s, "Prepared from regenerated slide assets, no initialized-minus-z no-motor row shown", Inches(0.75), Inches(6.35), Inches(12.0), Inches(0.25), size=11, color="muted")
    add_footer(s, slide_no)
    notes.append(("Title", ["Use this to frame the specific gap: collaborators have not seen the initialized-minus-z simulations yet."]))
    slide_no += 1

    s = blank(prs)
    add_title(s, "Motivation", "Why this set is useful before discussing membrane or motor-cluster variability")
    add_bullets(
        s,
        [
            "The random-initialization simulations mix two effects: polarity generation and motor-driven transport.",
            "The initialized-minus-z set isolates what happens when axial polarity is already present.",
            "This makes the movie/heatmap story easier: we can ask how motor organization amplifies or preserves a supplied polarity.",
            "The comparison also gives a baseline for later hypotheses about membrane deformation or mobile motor clusters.",
        ],
        Inches(0.85),
        Inches(1.55),
        Inches(11.8),
        Inches(4.6),
        size=23,
    )
    add_footer(s, slide_no)
    notes.append(("Motivation", ["Emphasize that this is not claiming the biological system starts perfectly aligned; it is a controlled simulation test."]))
    slide_no += 1

    s = blank(prs)
    add_title(s, "Simulation Design", "Same annular geometry; actin initialized along minus-z")
    add_card(
        s,
        "Geometry and filaments",
        [
            "Annular cylinder domain",
            "256 actin filaments",
            "Initial filament direction: 0 0 -1",
            "800 s analysis window",
        ],
        Inches(0.65),
        Inches(1.45),
        Inches(3.8),
        Inches(4.65),
        accent="teal",
    )
    add_card(
        s,
        "Motor conditions",
        [
            "Inner-wall myosin clusters",
            "Rotatable motors",
            "Fixed-orientation motors",
            "Cluster counts: 10, 20, 40, 60, 80",
        ],
        Inches(4.75),
        Inches(1.45),
        Inches(3.8),
        Inches(4.65),
        accent="orange",
    )
    add_card(
        s,
        "Parameter sweeps",
        [
            "Fixed total motors = 480",
            "12 motors per cluster",
            "Crosslinker ratios: 1:4, 1:8, 1:16",
            "10 runs per ratio where available",
        ],
        Inches(8.85),
        Inches(1.45),
        Inches(3.8),
        Inches(4.65),
        accent="purple",
    )
    add_textbox(
        s,
        "No-motor panels are omitted here because the available no-motor runs use random initial filament orientations, not initialized minus-z.",
        Inches(0.75),
        Inches(6.40),
        Inches(11.9),
        Inches(0.35),
        size=12,
        color="muted",
    )
    add_footer(s, slide_no)
    notes.append(("Simulation design", ["This is the slide to mention the control issue explicitly but briefly."]))
    slide_no += 1

    s = blank(prs)
    add_title(s, "How To Read The Figures", "Unwrapped heatmaps carry the spatial story; metrics summarize the sweep")
    add_card(
        s,
        "Unwrapped annulus",
        [
            "x-axis: circumference after unwrapping the annulus",
            "y-axis: z position",
            "columns: time through the 800 s window",
            "rows: cluster count and motor model",
        ],
        Inches(0.65),
        Inches(1.42),
        Inches(4.05),
        Inches(4.95),
        accent="teal",
    )
    add_card(
        s,
        "Heatmap signal",
        [
            "Color reports normalized areal density",
            "Bright vertical features indicate axial bundles",
            "Late-time concentration highlights motor-driven reorganization",
        ],
        Inches(4.95),
        Inches(1.42),
        Inches(3.85),
        Inches(4.95),
        accent="orange",
    )
    add_card(
        s,
        "Metrics",
        [
            "Mean |v_z|: axial transport speed",
            "Active transport fraction: persistent motion above threshold",
            "Directionality index: |Delta z| / AUC(|v_z|)",
            "AUC(|v_z|): total axial motion",
        ],
        Inches(9.05),
        Inches(1.42),
        Inches(3.6),
        Inches(4.95),
        accent="purple",
    )
    add_footer(s, slide_no)
    notes.append(("Reading guide", ["Use this before the plots so the audience does not have to decode axes while interpreting the result."]))
    slide_no += 1

    s = blank(prs)
    add_title(s, "Representative Movie Slots", "Use these as the first visual comparison, then transition to heatmaps")
    add_label(s, "Fixed total motors = 480; 80 clusters", Inches(0.8), Inches(1.35), Inches(5.7), color="teal")
    add_image(s, "unwrapped_heatmaps/total480_1to16_high_counts_timeline_heatmap.png", Inches(0.55), Inches(1.76), Inches(6.2), Inches(4.7))
    add_label(s, "12 motors per cluster; 80 clusters", Inches(6.95), Inches(1.35), Inches(5.7), color="orange")
    add_image(s, "unwrapped_heatmaps/mpc12_1to16_high_counts_timeline_heatmap.png", Inches(6.70), Inches(1.76), Inches(6.2), Inches(4.7))
    add_textbox(
        s,
        "Movie files were not present in the workspace, so this slide uses the matched heatmap summaries as placeholders.",
        Inches(0.75),
        Inches(6.55),
        Inches(11.9),
        Inches(0.3),
        size=11,
        color="muted",
    )
    add_footer(s, slide_no)
    notes.append(("Movie slots", ["Drop in the actual movies when available. Use the heatmaps if presenting from this generated deck as-is."]))
    slide_no += 1

    s = blank(prs)
    add_title(s, "Spatial Reorganization: Fixed Total Motor Budget", "Total motors = 480, xlink ratio 1:16")
    add_image(s, "unwrapped_heatmaps/total480_1to16_low_counts_timeline_heatmap.png", Inches(0.45), Inches(1.18), Inches(12.4), Inches(5.65))
    add_footer(s, slide_no)
    notes.append(("Fixed total motor budget", ["At fixed total motors, adding more clusters also reduces motors per cluster, so the trend is not a pure total-motor increase."]))
    slide_no += 1

    s = blank(prs)
    add_title(s, "Spatial Reorganization: Increasing Total Motors", "12 motors per cluster, xlink ratio 1:16")
    add_image(s, "unwrapped_heatmaps/mpc12_1to16_high_counts_timeline_heatmap.png", Inches(0.48), Inches(1.18), Inches(12.3), Inches(5.65))
    add_footer(s, slide_no)
    notes.append(("Increasing total motors", ["This is the cleanest visual for high cluster counts because total motor number rises with cluster count."]))
    slide_no += 1

    s = blank(prs)
    add_title(s, "Mean Axial Speed", "Rotatable motors produce faster axial motion across the initialized-minus-z sweep")
    add_image(s, "transport_metrics/mean_vz_abs.png", Inches(0.55), Inches(1.15), Inches(12.25), Inches(5.8))
    add_footer(s, slide_no)
    notes.append(("Mean speed", ["Point out that the 12-motors-per-cluster sweep climbs strongly with cluster count; fixed-total-motor cases are flatter."]))
    slide_no += 1

    s = blank(prs)
    add_title(s, "Active Fraction And Directionality", "Transport becomes more persistent and more coherent when motors can rotate")
    add_image(s, "transport_metrics/active_transport_fraction.png", Inches(0.30), Inches(1.15), Inches(6.35), Inches(5.65))
    add_image(s, "transport_metrics/directionality_index.png", Inches(6.68), Inches(1.15), Inches(6.35), Inches(5.65))
    add_footer(s, slide_no)
    notes.append(("Active fraction and directionality", ["These metrics separate 'moving a lot' from 'moving coherently in one axial direction'."]))
    slide_no += 1

    s = blank(prs)
    add_title(s, "Random vs Initialized Minus-Z", "Supplying initial polarity shifts both speed and coherence")
    add_image(s, "four_way_comparison/mean_vz_abs.png", Inches(0.35), Inches(1.15), Inches(6.25), Inches(5.65))
    add_image(s, "four_way_comparison/directionality_index.png", Inches(6.75), Inches(1.15), Inches(6.25), Inches(5.65))
    add_footer(s, slide_no)
    notes.append(("Random vs aligned", ["Use this slide to make the strongest comparison: same model categories, different initial filament organization."]))
    slide_no += 1

    s = blank(prs)
    add_title(s, "Main Interpretation", "The initialized-minus-z condition reveals a high-transport regime")
    add_bullets(
        s,
        [
            "Initial filament polarity is enough to unlock coherent axial transport in the annular geometry.",
            "Rotatable motors are consistently stronger than fixed-orientation motors in speed, active fraction, AUC, and directionality.",
            "When motors per cluster is fixed, increasing cluster count also increases total motor number and gives the clearest monotonic transport increase.",
            "At fixed total motors, the response is flatter because cluster number and motors per cluster trade off against each other.",
        ],
        Inches(0.85),
        Inches(1.45),
        Inches(11.7),
        Inches(4.9),
        size=22,
    )
    add_footer(s, slide_no)
    notes.append(("Interpretation", ["Keep this slide as the concise verbal summary after the data-heavy figures."]))
    slide_no += 1

    s = blank(prs)
    add_title(s, "Discussion Points", "What this motivates next")
    add_card(
        s,
        "For this deck",
        [
            "Do not show random-initialized no-motor controls as initialized-minus-z controls.",
            "Use the 10-cluster cases as the lowest motorized visual reference.",
            "Mention active-fraction thresholds only as inherited analysis thresholds.",
        ],
        Inches(0.65),
        Inches(1.45),
        Inches(5.75),
        Inches(4.8),
        accent="teal",
    )
    add_card(
        s,
        "Next simulations",
        [
            "Matched initialized-minus-z no-motor controls only if a formal baseline is required.",
            "Test imperfect or deformable inner wall geometry for membrane-coupled cluster motion.",
            "Allow motor clusters to redistribute to test whether clustering can emerge dynamically.",
        ],
        Inches(6.9),
        Inches(1.45),
        Inches(5.75),
        Inches(4.8),
        accent="orange",
    )
    add_footer(s, slide_no)
    notes.append(("Discussion", ["This slide connects the result back to the collaborator observation and the next model variants."]))

    OUT.parent.mkdir(parents=True, exist_ok=True)
    prs.save(OUT)

    lines = ["# Initially aligned minus-z slide notes", ""]
    for idx, (title, slide_notes) in enumerate(notes, start=1):
        lines.append(f"## Slide {idx}: {title}")
        for note in slide_notes:
            lines.append(f"- {note}")
        lines.append("")
    NOTES.write_text("\n".join(lines), encoding="utf-8")


if __name__ == "__main__":
    build_deck()
    print(OUT)
    print(NOTES)
